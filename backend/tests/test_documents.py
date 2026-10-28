import io
import docx
import pypdf
from fastapi.testclient import TestClient

def test_health_check(client: TestClient):
    response = client.get("/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "healthy"
    assert data["service"] == "AI Document Intelligence System"
    assert ".pdf" in data["supported_formats"]
    assert ".docx" in data["supported_formats"]
    assert ".txt" in data["supported_formats"]

def test_root_endpoint(client: TestClient):
    response = client.get("/")
    assert response.status_code == 200
    data = response.json()
    assert "version" in data
    assert "Phase" in data["phase"]

def test_supported_types(client: TestClient):
    response = client.get("/api/documents/supported-types")
    assert response.status_code == 200
    data = response.json()
    assert ".pdf" in data["allowed_extensions"]
    assert data["max_upload_size_mb"] >= 10

def test_upload_empty_file_rejected(client: TestClient):
    files = {"file": ("empty.txt", b"", "text/plain")}
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()

def test_upload_unsupported_file_extension(client: TestClient):
    files = {"file": ("malicious.exe", b"MZ\x90\x00", "application/octet-stream")}
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 400
    assert "unsupported" in response.json()["detail"].lower()

def test_upload_corrupt_pdf_rejected(client: TestClient):
    files = {"file": ("corrupt.pdf", b"Not a valid PDF header", "application/pdf")}
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 400
    assert "invalid pdf" in response.json()["detail"].lower()

def test_upload_corrupt_docx_rejected(client: TestClient):
    files = {"file": ("corrupt.docx", b"PKFakeZipWithoutDocumentXml", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 400
    assert "docx" in response.json()["detail"].lower()

def test_upload_txt_file_with_page_breaks(client: TestClient):
    content = "First page content with important notes.\x0cSecond page content with intelligence data."
    files = {"file": ("notes.txt", content.encode("utf-8"), "text/plain")}
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert data["filename"] == "notes.txt"
    assert data["file_type"] == "TXT"
    assert data["total_pages"] == 2
    assert data["extracted_pages"][0]["page_number"] == 1
    assert "First page content" in data["extracted_pages"][0]["text"]
    assert data["extracted_pages"][1]["page_number"] == 2
    assert "Second page content" in data["extracted_pages"][1]["text"]

def test_upload_docx_file_with_page_break(client: TestClient):
    doc = docx.Document()
    doc.add_heading("Project Overview", level=1)
    doc.add_paragraph("This is paragraph one of document intelligence system.")
    doc.add_page_break()
    doc.add_paragraph("This is paragraph two on the second page.")
    
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    files = {"file": ("overview.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert data["filename"] == "overview.docx"
    assert data["file_type"] == "DOCX"
    assert data["total_pages"] >= 2
    assert "Project Overview" in data["extracted_pages"][0]["text"]
    assert data["total_words"] > 0

def test_upload_pdf_file(client: TestClient):
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.add_blank_page(width=300, height=300)
    
    buf = io.BytesIO()
    writer.write(buf)
    pdf_bytes = buf.getvalue()

    files = {"file": ("sample.pdf", pdf_bytes, "application/pdf")}
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert data["file_type"] == "PDF"
    assert data["total_pages"] == 2
    assert len(data["extracted_pages"]) == 2
    assert data["extracted_pages"][0]["page_number"] == 1
    assert data["extracted_pages"][1]["page_number"] == 2

def test_validate_endpoint_valid_and_invalid(client: TestClient):
    # Valid text
    files = {"file": ("sample.txt", b"Valid text file content", "text/plain")}
    response = client.post("/api/documents/validate", files=files)
    assert response.status_code == 200
    assert response.json()["is_valid"] is True
    
    # Invalid extension
    files_inv = {"file": ("test.exe", b"executable payload", "application/octet-stream")}
    response_inv = client.post("/api/documents/validate", files=files_inv)
    assert response_inv.status_code == 200
    assert response_inv.json()["is_valid"] is False
    assert "unsupported" in response_inv.json()["error"].lower()

def test_sanitize_path_traversal(client: TestClient):
    content = b"Safe content"
    files = {"file": ("../../../../etc/passwd.txt", content, "text/plain")}
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 200
    data = response.json()
    assert ".." not in data["filename"]
    assert "/" not in data["filename"]
    assert "\\" not in data["filename"]

def test_security_headers_present(client: TestClient):
    response = client.get("/health")
    assert response.status_code == 200
    assert response.headers.get("X-Content-Type-Options") == "nosniff"
    assert response.headers.get("X-Frame-Options") == "DENY"
    assert response.headers.get("Referrer-Policy") == "strict-origin-when-cross-origin"
    assert "X-Process-Time" in response.headers

def test_upload_oversized_file_rejected(client: TestClient, monkeypatch):
    # Temporarily lower MAX_UPLOAD_SIZE_BYTES to 100KB to test 413 without allocating 16MB in tests
    from app.config import settings
    monkeypatch.setattr(settings, "MAX_UPLOAD_SIZE_BYTES", 100 * 1024)
    monkeypatch.setattr(settings, "MAX_UPLOAD_SIZE_MB", 1)

    oversized_data = b"A" * (120 * 1024)
    files = {"file": ("large.txt", oversized_data, "text/plain")}
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 413
    assert "exceeds maximum allowed size" in response.json()["detail"].lower()

def test_txt_with_null_bytes_rejected(client: TestClient):
    binary_payload = b"Hello\x00World\x00BinaryData"
    files = {"file": ("payload.txt", binary_payload, "text/plain")}
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 400
    assert "binary data" in response.json()["detail"].lower()

def test_page_boundary_types_and_approximation_flags(client: TestClient):
    # 1. PDF physical page
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    pdf_bytes = buf.getvalue()
    res_pdf = client.post("/api/documents/upload", files={"file": ("report.pdf", pdf_bytes, "application/pdf")})
    assert res_pdf.status_code == 200
    assert res_pdf.json()["extracted_pages"][0]["boundary_type"] == "physical"
    assert res_pdf.json()["extracted_pages"][0]["is_approximate"] is False

    # 2. TXT with form feed
    txt_ff = "Page1\x0cPage2"
    res_ff = client.post("/api/documents/upload", files={"file": ("ff.txt", txt_ff.encode("utf-8"), "text/plain")})
    assert res_ff.status_code == 200
    assert res_ff.json()["extracted_pages"][0]["boundary_type"] == "form_feed"
    assert res_ff.json()["extracted_pages"][0]["is_approximate"] is False

    # 3. TXT without form feed (approximate / single)
    txt_simple = "Single continuous text block."
    res_simple = client.post("/api/documents/upload", files={"file": ("simple.txt", txt_simple.encode("utf-8"), "text/plain")})
    assert res_simple.status_code == 200
    assert res_simple.json()["extracted_pages"][0]["boundary_type"] == "approximate_chunk"
    assert res_simple.json()["extracted_pages"][0]["is_approximate"] is True

def test_sample_document_identification(client: TestClient):
    content = b"Demo briefing content"
    files = {"file": ("[Sample] Quarterly Briefing.txt", content, "text/plain")}
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 200
    assert response.json()["is_sample"] is True

    # Real user document
    files_user = {"file": ("user_tax_return.txt", content, "text/plain")}
    response_user = client.post("/api/documents/upload", files=files_user)
    assert response_user.status_code == 200
    assert response_user.json()["is_sample"] is False
